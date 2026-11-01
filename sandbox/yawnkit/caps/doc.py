"""doc 包：Word / Excel / Markdown / HTML / JSON / YAML / CSV 之间的读取与转换。"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from . import CapError, cap, clip_list, clip_text, dest, human, read_text, ref, src


def _stem(path: str) -> str:
    return Path(str(path).removeprefix("sandbox:")).stem


@cap("doc.docx_text")
def docx_text(path: str, max_chars: int = 6000) -> dict[str, Any]:
    import docx

    document = docx.Document(str(src(path, kind="Word 文档")))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    payload = clip_text("\n".join(parts), max(500, min(int(max_chars), 20000)))
    payload["paragraphs"] = len(document.paragraphs)
    payload["tables"] = len(document.tables)
    return payload


@cap("doc.docx_write")
def docx_write(text: str, output: str = "", title: str = "") -> dict[str, Any]:
    import docx

    document = docx.Document()
    if title:
        document.add_heading(str(title), level=1)
    for block in str(text).split("\n"):
        stripped = block.strip()
        if stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped:
            document.add_paragraph(stripped)
    target = dest(output, f"{title or '文档'}.docx")
    document.save(str(target))
    return {"file": ref(target), "size": human(target.stat().st_size)}


@cap("doc.xlsx_sheets")
def xlsx_sheets(path: str) -> dict[str, Any]:
    from openpyxl import load_workbook

    book = load_workbook(str(src(path, kind="Excel")), read_only=True, data_only=True)
    rows = [
        {"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column}
        for ws in book.worksheets
    ]
    book.close()
    return {"sheets": rows}


@cap("doc.xlsx_to_csv")
def xlsx_to_csv(path: str, sheet: str = "", output: str = "") -> dict[str, Any]:
    import pandas as pd

    frame = pd.read_excel(src(path, kind="Excel"), sheet_name=sheet or 0)
    target = dest(output, f"{_stem(path)}{'_' + sheet if sheet else ''}.csv")
    frame.to_csv(target, index=False, encoding="utf-8-sig")
    return {
        "file": ref(target),
        "rows": len(frame),
        "columns": list(map(str, frame.columns))[:30],
    }


@cap("doc.csv_to_xlsx")
def csv_to_xlsx(path: str, output: str = "") -> dict[str, Any]:
    from .data import load_table

    frame = load_table(path)
    target = dest(output, f"{_stem(path)}.xlsx")
    frame.to_excel(target, index=False)
    return {"file": ref(target), "rows": len(frame)}


@cap("doc.markdown_to_html")
def markdown_to_html(
    path: str = "", text: str = "", output: str = ""
) -> dict[str, Any]:
    import markdown

    source = read_text(src(path)) if path else str(text)
    if not source.strip():
        raise CapError("给 path 或 text")
    body = markdown.markdown(source, extensions=["tables", "fenced_code"])
    html = f"<!doctype html><meta charset='utf-8'><style>body{{font-family:sans-serif;max-width:760px;margin:2em auto;line-height:1.7}}table{{border-collapse:collapse}}td,th{{border:1px solid #ccc;padding:4px 8px}}</style>{body}"
    target = dest(output, f"{_stem(path) if path else '文档'}.html")
    target.write_text(html, encoding="utf-8")
    return {"file": ref(target)}


@cap("doc.html_to_text")
def html_to_text(path: str, max_chars: int = 6000) -> dict[str, Any]:
    from html.parser import HTMLParser

    class Strip(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.parts: list[str] = []
            self.skip = 0

        def handle_starttag(self, tag: str, attrs: Any) -> None:  # noqa: ARG002
            if tag in {"script", "style"}:
                self.skip += 1

        def handle_endtag(self, tag: str) -> None:
            if tag in {"script", "style"} and self.skip:
                self.skip -= 1

        def handle_data(self, data: str) -> None:
            if not self.skip and data.strip():
                self.parts.append(data.strip())

    parser = Strip()
    parser.feed(read_text(src(path)))
    return clip_text("\n".join(parser.parts), max(500, min(int(max_chars), 20000)))


@cap("doc.json_to_yaml")
def json_to_yaml(path: str, output: str = "") -> dict[str, Any]:
    import yaml

    data = json.loads(read_text(src(path)))
    target = dest(output, f"{_stem(path)}.yaml")
    target.write_text(
        yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    return {"file": ref(target)}


@cap("doc.yaml_to_json")
def yaml_to_json(path: str, output: str = "") -> dict[str, Any]:
    import yaml

    try:
        data = yaml.safe_load(read_text(src(path)))
    except yaml.YAMLError as exc:
        raise CapError(f"YAML 格式不对：{exc}") from exc
    target = dest(output, f"{_stem(path)}.json")
    target.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    return {"file": ref(target)}


@cap("doc.json_to_csv")
def json_to_csv(path: str, output: str = "") -> dict[str, Any]:
    import pandas as pd

    data = json.loads(read_text(src(path)))
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [data])
    frame = pd.json_normalize(data)
    target = dest(output, f"{_stem(path)}.csv")
    frame.to_csv(target, index=False, encoding="utf-8-sig")
    return {
        "file": ref(target),
        "rows": len(frame),
        "columns": list(map(str, frame.columns))[:30],
    }


@cap("doc.csv_to_json")
def csv_to_json(path: str, output: str = "") -> dict[str, Any]:
    from .data import load_table

    frame = load_table(path)
    target = dest(output, f"{_stem(path)}.json")
    frame.to_json(target, orient="records", force_ascii=False, indent=2)
    return {"file": ref(target), "rows": len(frame)}


@cap("doc.csv_preview")
def csv_preview(path: str, rows: int = 10) -> dict[str, Any]:
    text = read_text(src(path))
    reader = csv.reader(text.splitlines())
    items = [
        row
        for _, row in zip(range(max(1, min(int(rows), 50)) + 1), reader, strict=False)
    ]
    return clip_list(items)


@cap("doc.table_to_markdown")
def table_to_markdown(path: str, rows: int = 30) -> dict[str, Any]:
    from .data import load_table

    frame = load_table(path).head(max(1, min(int(rows), 200)))
    return clip_text(frame.to_markdown(index=False))
